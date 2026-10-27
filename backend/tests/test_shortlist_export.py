import io
import json
import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.core.security import create_access_token, hash_password
from app.main import app
from app.models.candidate import Candidate
from app.models.job import Job
from app.models.user import User
from app.services.export_service import ExportService, sanitize_excel_cell
from tests.test_auth import client, test_engine, TestingSessionLocal

def test_sanitize_excel_cell_formula_injection():
    assert sanitize_excel_cell("=1+1") == "'=1+1"
    assert sanitize_excel_cell("+SUM(A1:A10)") == "'+SUM(A1:A10)"
    assert sanitize_excel_cell("-100") == "'-100"
    assert sanitize_excel_cell("@SUM(B1)") == "'@SUM(B1)"
    assert sanitize_excel_cell("Normal Text") == "Normal Text"
    assert sanitize_excel_cell(None) == ""

def test_excel_export_generation_and_structure():
    candidates = [
        Candidate(
            id=1,
            job_id=10,
            recruiter_id=1,
            original_filename="jane_doe.pdf",
            file_type="pdf",
            file_size=1024,
            processing_status="completed",
            full_name="Jane Doe",
            email="jane@example.com",
            phone="555-1234",
            location="San Francisco, CA",
            relevance_score=88.5,
            fit_category="Good Fit",
            ranking_position=1,
            matched_skills="Python, PyTorch, FastSearch",
            linkedin_url="https://linkedin.com/in/janedoe"
        ),
        Candidate(
            id=2,
            job_id=10,
            recruiter_id=1,
            original_filename="malicious.pdf",
            file_type="pdf",
            file_size=2048,
            processing_status="completed",
            full_name="=CMD|' /C calc'!A0", # Injection attack in name
            email="+victim@domain.com",
            phone="-555-9999",
            location="New York, NY",
            relevance_score=52.0,
            fit_category="Potential Fit",
            ranking_position=2,
            matched_skills="@ExcelFormulaInjection",
            linkedin_url=None
        )
    ]

    xlsx_bytes = ExportService.generate_excel_export(
        candidates=candidates,
        job_title="Senior ML Engineer",
        limit_requested=100
    )
    assert len(xlsx_bytes) > 0

    # Load back with openpyxl and assert structure
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Ranked Shortlist" in wb.sheetnames
    ws = wb["Ranked Shortlist"]

    # Header check
    assert ws["A1"].value == "VERISIFT — CANDIDATE RANKING SHORTLIST"
    assert ws["A4"].value == "Rank"
    assert ws["D4"].value == "Candidate Name"

    # Row 5 (Jane Doe)
    assert ws["A5"].value == 1
    assert ws["B5"].value == "88.5"
    assert ws["C5"].value == "Good Fit"
    assert ws["D5"].value == "Jane Doe"
    assert ws["E5"].value == "jane@example.com"

    # Row 6 (Formula injection sanitization check)
    assert ws["A6"].value == 2
    assert ws["D6"].value == "'=CMD|' /C calc'!A0" # Safely escaped with leading apostrophe
    assert ws["E6"].value == "'+victim@domain.com"
    assert ws["F6"].value == "'-555-9999"
    assert ws["H6"].value == "'@ExcelFormulaInjection"

def test_pdf_export_generation_and_headers():
    candidates = [
        Candidate(
            id=1,
            job_id=10,
            recruiter_id=1,
            original_filename="alex_rivera.pdf",
            file_type="pdf",
            file_size=1024,
            processing_status="completed",
            full_name="Alex Rivera",
            email="alex@rivera.ai",
            phone="415-555-0199",
            location="San Francisco, CA",
            relevance_score=94.2,
            fit_category="Good Fit",
            ranking_position=1,
            matched_skills="Deep Learning, PyTorch, Transformers"
        )
    ]
    pdf_bytes = ExportService.generate_pdf_export(
        candidates=candidates,
        job_title="Lead AI Architect",
        limit_requested=100
    )
    assert len(pdf_bytes) > 0
    assert pdf_bytes.startswith(b"%PDF-")

@pytest.fixture
def export_test_data():
    db = TestingSessionLocal()
    try:
        # Recruiter 1
        u1 = db.query(User).filter(User.email == "export_rec1@verisift.io").first()
        if not u1:
            u1 = User(
                email="export_rec1@verisift.io",
                password_hash=hash_password("Password123!"),
                full_name="Export Recruiter One"
            )
            db.add(u1)
            db.commit()
            db.refresh(u1)

        # Recruiter 2
        u2 = db.query(User).filter(User.email == "export_rec2@verisift.io").first()
        if not u2:
            u2 = User(
                email="export_rec2@verisift.io",
                password_hash=hash_password("Password123!"),
                full_name="Export Recruiter Two"
            )
            db.add(u2)
            db.commit()
            db.refresh(u2)

        # Job for Recruiter 1
        job1 = db.query(Job).filter(Job.recruiter_id == u1.id, Job.title == "Export Test Role").first()
        if not job1:
            job1 = Job(
                recruiter_id=u1.id,
                title="Export Test Role",
                department="Engineering",
                description="Backend engineer role"
            )
            db.add(job1)
            db.commit()
            db.refresh(job1)

        # Add 3 completed ranked candidates
        db.query(Candidate).filter(Candidate.job_id == job1.id).delete()
        c1 = Candidate(
            job_id=job1.id,
            recruiter_id=u1.id,
            original_filename="cand1.pdf",
            stored_filename="c1.pdf",
            file_type="pdf",
            file_size=100,
            file_hash="hash1",
            processing_status="completed",
            full_name="Alice Candidate",
            email="alice@test.com",
            relevance_score=92.5,
            fit_category="Good Fit",
            ranking_position=1
        )
        c2 = Candidate(
            job_id=job1.id,
            recruiter_id=u1.id,
            original_filename="cand2.pdf",
            stored_filename="c2.pdf",
            file_type="pdf",
            file_size=100,
            file_hash="hash2",
            processing_status="completed",
            full_name="Bob Candidate",
            email="bob@test.com",
            relevance_score=68.0,
            fit_category="Potential Fit",
            ranking_position=2
        )
        c3 = Candidate(
            job_id=job1.id,
            recruiter_id=u1.id,
            original_filename="cand3.pdf",
            stored_filename="c3.pdf",
            file_type="pdf",
            file_size=100,
            file_hash="hash3",
            processing_status="completed",
            full_name="Charlie Candidate",
            email="charlie@test.com",
            relevance_score=21.0,
            fit_category="No Fit",
            ranking_position=3
        )
        db.add_all([c1, c2, c3])
        db.commit()

        # Job with NO candidates for testing 400
        empty_job = db.query(Job).filter(Job.recruiter_id == u1.id, Job.title == "Empty Export Role").first()
        if not empty_job:
            empty_job = Job(
                recruiter_id=u1.id,
                title="Empty Export Role",
                department="QA",
                description="No applicants role"
            )
            db.add(empty_job)
            db.commit()
            db.refresh(empty_job)

        token1 = create_access_token({"sub": str(u1.id), "email": u1.email})
        token2 = create_access_token({"sub": str(u2.id), "email": u2.email})

        yield {
            "u1_headers": {"Authorization": f"Bearer {token1}"},
            "u2_headers": {"Authorization": f"Bearer {token2}"},
            "job_id": job1.id,
            "empty_job_id": empty_job.id,
        }
    finally:
        db.close()

def test_api_export_shortlist_excel_top100(export_test_data):
    headers = export_test_data["u1_headers"]
    job_id = export_test_data["job_id"]

    res = client.get(f"/api/jobs/{job_id}/candidates/export?limit=100&format=xlsx", headers=headers)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in res.headers["content-disposition"]
    assert res.headers["content-disposition"].endswith('.xlsx"')

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Ranked Shortlist"]
    assert ws["D5"].value == "Alice Candidate"
    assert ws["D6"].value == "Bob Candidate"
    assert ws["D7"].value == "Charlie Candidate"

def test_api_export_shortlist_pdf_top200(export_test_data):
    headers = export_test_data["u1_headers"]
    job_id = export_test_data["job_id"]

    res = client.get(f"/api/jobs/{job_id}/candidates/export?limit=200&format=pdf", headers=headers)
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    assert "attachment; filename=" in res.headers["content-disposition"]
    assert res.headers["content-disposition"].endswith('.pdf"')
    assert res.content.startswith(b"%PDF-")

def test_api_export_shortlist_custom_limit(export_test_data):
    headers = export_test_data["u1_headers"]
    job_id = export_test_data["job_id"]

    # Limit to top 2 candidates
    res = client.get(f"/api/jobs/{job_id}/candidates/export?limit=2&format=xlsx", headers=headers)
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Ranked Shortlist"]
    # Row 5 and 6 should be populated; row 7 should be None
    assert ws["D5"].value == "Alice Candidate"
    assert ws["D6"].value == "Bob Candidate"
    assert ws["D7"].value is None

def test_api_export_empty_job_returns_400(export_test_data):
    headers = export_test_data["u1_headers"]
    empty_job_id = export_test_data["empty_job_id"]

    res = client.get(f"/api/jobs/{empty_job_id}/candidates/export?limit=100&format=xlsx", headers=headers)
    assert res.status_code == 400
    assert "no completed candidates" in res.json()["detail"].lower()

def test_api_export_unauthorized_other_recruiter(export_test_data):
    # Recruiter 2 trying to export Recruiter 1's job
    headers2 = export_test_data["u2_headers"]
    job_id = export_test_data["job_id"]

    res = client.get(f"/api/jobs/{job_id}/candidates/export?limit=100&format=xlsx", headers=headers2)
    assert res.status_code == 404

def test_api_export_missing_auth_header(export_test_data):
    job_id = export_test_data["job_id"]
    res = client.get(f"/api/jobs/{job_id}/candidates/export?limit=100&format=xlsx")
    assert res.status_code == 401

def test_api_export_invalid_format(export_test_data):
    headers = export_test_data["u1_headers"]
    job_id = export_test_data["job_id"]

    res = client.get(f"/api/jobs/{job_id}/candidates/export?limit=100&format=csv", headers=headers)
    assert res.status_code == 422 # FastAPI query validation error

def test_api_export_invalid_limit_bounds(export_test_data):
    headers = export_test_data["u1_headers"]
    job_id = export_test_data["job_id"]

    # Limit = 0
    res0 = client.get(f"/api/jobs/{job_id}/candidates/export?limit=0&format=xlsx", headers=headers)
    assert res0.status_code == 422

    # Limit = -10
    res_neg = client.get(f"/api/jobs/{job_id}/candidates/export?limit=-10&format=xlsx", headers=headers)
    assert res_neg.status_code == 422

    # Limit > 1000
    res_large = client.get(f"/api/jobs/{job_id}/candidates/export?limit=1500&format=xlsx", headers=headers)
    assert res_large.status_code == 422

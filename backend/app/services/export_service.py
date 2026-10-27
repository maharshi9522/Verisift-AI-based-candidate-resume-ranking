import io
import re
from datetime import datetime, timezone
from typing import List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch, mm
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.models.candidate import Candidate


def sanitize_excel_cell(val: Optional[str]) -> str:
    """
    Prevents CSV / Excel formula injection.
    Prefixes values starting with '=', '+', '-', or '@' with a single apostrophe.
    """
    if val is None:
        return ""
    val_str = str(val).strip()
    if val_str and val_str[0] in ("=", "+", "-", "@"):
        return f"'{val_str}"
    return val_str


class NumberedCanvas(canvas.Canvas):
    """
    Canvas that performs a two-pass calculation of total pages,
    adding professional running headers and 'Page X of Y' footers.
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B")) # Slate 500

        # Header (pages 2+)
        if self._pageNumber > 1:
            self.drawString(36, 805, "Verisift — Candidate Ranking Shortlist")
            self.drawRightString(559, 805, datetime.now(timezone.utc).strftime("%b %d, %Y"))
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(36, 798, 559, 798)

        # Footer (all pages)
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(36, 40, 559, 40)
        self.drawString(36, 28, "Confidential — For Authorized Recruiter Use Only")
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(559, 28, page_str)
        self.restoreState()


class ExportService:
    @staticmethod
    def generate_excel_export(
        candidates: List[Candidate],
        job_title: str,
        limit_requested: int
    ) -> bytes:
        """
        Generates a styled, filterable Excel (.xlsx) workbook for the candidate shortlist.
        Features:
        - Metadata block with job title, export timestamp, candidate count.
        - Frozen header row.
        - Auto-filters enabled.
        - Formula injection sanitization.
        - Professional color palette and typography.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ranked Shortlist"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        navy_dark = "1E293B" # Slate 800
        blue_accent = "2563EB" # Blue 600
        light_gray = "F8FAFC"
        border_color = "CBD5E1"

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
        title_font = Font(name="Calibri", size=16, bold=True, color=blue_accent)
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
        regular_font = Font(name="Calibri", size=10, color="0F172A")
        bold_font = Font(name="Calibri", size=10, bold=True, color="0F172A")

        thin_side = Side(border_style="thin", color=border_color)
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        # Metadata Header
        ws["A1"] = "VERISIFT — CANDIDATE RANKING SHORTLIST"
        ws["A1"].font = title_font
        ws["A2"] = f"Job Role: {job_title} | Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} | Requested: Top {limit_requested} | Available: {len(candidates)}"
        ws["A2"].font = subtitle_font

        # Blank row at 3
        # Table Headers at Row 4
        headers = [
            ("Rank", 8, center_align),
            ("Score / 100", 14, center_align),
            ("Predicted Fit", 15, center_align),
            ("Candidate Name", 24, left_align),
            ("Email", 28, left_align),
            ("Phone", 18, left_align),
            ("Location", 20, left_align),
            ("Matched Skills / Evidence", 30, left_align),
            ("Strongest Match", 32, left_align),
            ("Top Missing Requirement", 32, left_align),
            ("Peer Comparison", 35, left_align),
            ("Resume File", 25, left_align),
            ("LinkedIn", 25, left_align),
            ("GitHub", 25, left_align),
            ("Portfolio", 25, left_align),
        ]

        header_row = 4
        for col_idx, (header_text, _, alignment) in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=header_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = alignment
            cell.border = cell_border

        # Populate candidate rows
        for row_idx, cand in enumerate(candidates, start=5):
            # Display name fallback
            disp_name = cand.full_name or cand.original_filename.rsplit(".", 1)[0].replace("_", " ")
            rank_val = cand.ranking_position if cand.ranking_position is not None else (row_idx - 4)
            score_val = f"{cand.relevance_score:.1f}" if cand.relevance_score is not None else "N/A"
            fit_val = cand.fit_category or "Unranked"

            # Parse explanation if available
            strongest_match = ""
            top_missing = ""
            peer_comp = ""
            if cand.explanation:
                try:
                    import json
                    expl_data = json.loads(cand.explanation)
                    strongest_match = expl_data.get("strongest_match", "")
                    top_missing = expl_data.get("top_missing_requirement") or "None"
                    peer_comp = expl_data.get("peer_comparison", "")
                except Exception:
                    pass

            # Strip formula injection
            row_data = [
                rank_val,
                score_val,
                fit_val,
                sanitize_excel_cell(disp_name),
                sanitize_excel_cell(cand.email),
                sanitize_excel_cell(cand.phone),
                sanitize_excel_cell(cand.location),
                sanitize_excel_cell(cand.matched_skills),
                sanitize_excel_cell(strongest_match),
                sanitize_excel_cell(top_missing),
                sanitize_excel_cell(peer_comp),
                sanitize_excel_cell(cand.original_filename),
                sanitize_excel_cell(cand.linkedin_url),
                sanitize_excel_cell(cand.github_url),
                sanitize_excel_cell(cand.portfolio_url),
            ]

            is_even = (row_idx % 2 == 0)
            row_fill = PatternFill(start_color="F1F5F9" if is_even else "FFFFFF", fill_type="solid")

            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = bold_font if col_idx in (1, 2) else regular_font
                cell.fill = row_fill
                cell.border = cell_border
                cell.alignment = headers[col_idx - 1][2]

        # Column widths
        for col_idx, (_, width, _) in enumerate(headers, start=1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        # Freeze panes at row 5 (below headers)
        ws.freeze_panes = "A5"

        # Enable auto-filter
        if len(candidates) > 0:
            last_col_letter = get_column_letter(len(headers))
            last_row = 4 + len(candidates)
            ws.auto_filter.ref = f"A4:{last_col_letter}{last_row}"

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def generate_pdf_export(
        candidates: List[Candidate],
        job_title: str,
        limit_requested: int
    ) -> bytes:
        """
        Generates a recruiter-friendly, multi-page PDF shortlist using ReportLab.
        Features:
        - Professional header with branding, job title, date, count.
        - Table headers repeated automatically on every page.
        - Two-pass page numbering ("Page X of Y").
        - Clean typography, cell wrapping, and alternating row colors.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=36,
            rightMargin=36,
            topMargin=45,
            bottomMargin=45
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#1E293B")
        )
        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#64748B")
        )
        col_hdr_style = ParagraphStyle(
            "ColHeader",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.white,
            alignment=1 # Center
        )
        cell_style = ParagraphStyle(
            "CellNormal",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9.5,
            textColor=colors.HexColor("#0F172A")
        )
        cell_bold_center = ParagraphStyle(
            "CellBoldCenter",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#1E293B"),
            alignment=1
        )
        cell_score = ParagraphStyle(
            "CellScore",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#2563EB"), # Blue
            alignment=1
        )

        story = []

        # Top Header Card
        story.append(Paragraph("<b>VERISIFT</b> — Candidate Ranking Shortlist", title_style))
        story.append(Spacer(1, 4))
        meta_text = (
            f"<b>Job Requisition:</b> {job_title} &nbsp;|&nbsp; "
            f"<b>Export Date:</b> {datetime.now(timezone.utc).strftime('%B %d, %Y %H:%M UTC')}<br/>"
            f"<b>Scope:</b> Showing top {len(candidates)} of {len(candidates)} ranked candidates "
            f"(Requested: Top {limit_requested})"
        )
        story.append(Paragraph(meta_text, subtitle_style))
        story.append(Spacer(1, 10))

        # Table Layout: Total available width = 595.27 - 72 = 523.27 pt
        # Columns: Rank (32), Score (42), Fit (56), Candidate Details (150), Contact & Links (115), Skills/Evidence (128)
        # Sum = 32 + 42 + 56 + 150 + 115 + 128 = 523
        col_widths = [32, 42, 56, 150, 115, 128]

        table_data = [[
            Paragraph("Rank", col_hdr_style),
            Paragraph("Score", col_hdr_style),
            Paragraph("Fit Class", col_hdr_style),
            Paragraph("Candidate / Resume", col_hdr_style),
            Paragraph("Contact & Links", col_hdr_style),
            Paragraph("Matched Evidence / Skills", col_hdr_style),
        ]]

        for idx, cand in enumerate(candidates, start=1):
            rank_val = str(cand.ranking_position if cand.ranking_position is not None else idx)
            score_val = f"{cand.relevance_score:.1f}" if cand.relevance_score is not None else "—"
            fit_val = cand.fit_category or "Unranked"

            disp_name = cand.full_name or cand.original_filename.rsplit(".", 1)[0].replace("_", " ")
            cand_block = f"<b>{disp_name}</b><br/><font color='#64748B'>{cand.original_filename}</font>"

            # Contact info
            contact_parts = []
            if cand.email:
                contact_parts.append(cand.email)
            if cand.phone:
                contact_parts.append(cand.phone)
            if cand.location:
                contact_parts.append(cand.location)
            if cand.linkedin_url:
                contact_parts.append(f"<font color='#2563EB'>LinkedIn</font>")
            if cand.github_url:
                contact_parts.append(f"<font color='#2563EB'>GitHub</font>")

            contact_text = "<br/>".join(contact_parts) if contact_parts else "<font color='#94A3B8'>Not provided</font>"

            # Skills & Evidence explanation
            evidence_lines = []
            if cand.explanation:
                try:
                    import json
                    expl_data = json.loads(cand.explanation)
                    s_match = expl_data.get("strongest_match")
                    t_miss = expl_data.get("top_missing_requirement")
                    if s_match:
                        evidence_lines.append(f"<font color='#16A34A'><b>Match:</b></font> {s_match}")
                    if t_miss:
                        evidence_lines.append(f"<font color='#DC2626'><b>Gap:</b></font> {t_miss}")
                except Exception:
                    pass

            if not evidence_lines:
                if cand.matched_skills:
                    evidence_lines.append(f"<b>Skills:</b> {cand.matched_skills}")
                else:
                    evidence_lines.append("<font color='#94A3B8'>Semantic relevance match</font>")

            skills_text = "<br/>".join(evidence_lines)

            # Fit badge coloring
            fit_color = "#16A34A" if fit_val == "Good Fit" else ("#CA8A04" if fit_val == "Potential Fit" else "#DC2626")
            fit_html = f"<b><font color='{fit_color}'>{fit_val}</font></b>"

            table_data.append([
                Paragraph(f"#{rank_val}", cell_bold_center),
                Paragraph(score_val, cell_score),
                Paragraph(fit_html, cell_bold_center),
                Paragraph(cand_block, cell_style),
                Paragraph(contact_text, cell_style),
                Paragraph(skills_text, cell_style),
            ])

        pdf_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        pdf_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ]))

        story.append(pdf_table)

        doc.build(story, canvasmaker=NumberedCanvas)
        return buffer.getvalue()
